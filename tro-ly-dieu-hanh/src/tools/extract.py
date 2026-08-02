"""Trích xuất text từ tài liệu, có cache và có dò dấu hiệu injection.

Cache đánh khoá theo (file_id, modified_time). File Drive không đổi thì khoá
không đổi, nên lần hỏi thứ hai về cùng tài liệu tốn 0 lần gọi mạng và 0 lần
parse. File sửa rồi thì modified_time đổi → khoá đổi → tự động nạp lại. Không
cần cơ chế hết hạn, và cũng không có nguy cơ dùng nhầm nội dung cũ.

Việc dò injection chạy ngay lúc trích xuất và cảnh báo được lưu cùng cache, để
mọi chỗ dùng lại nội dung đều thấy cảnh báo mà không phải quét lại.
"""

from __future__ import annotations

import hashlib
import io
import json
import re
import unicodedata
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Protocol, Sequence

from ..core.guards import InjectionHit, scan_for_injection

# Trần ký tự cho một tài liệu. Vượt quá thì cắt và đánh dấu — một file 500 trang
# đưa nguyên vào prompt sẽ ngốn hết ngân sách context mà phần dùng đến chỉ vài đoạn.
MAX_CHARS = 40_000

GOOGLE_EXPORT_MIMES = {
    "application/vnd.google-apps.document": "text/plain",
    "application/vnd.google-apps.spreadsheet":
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/vnd.google-apps.presentation": "text/plain",
}


class ExtractionError(Exception):
    """Không trích xuất được. Không phải lỗi chí mạng — S5 quyết định có bỏ qua được không."""


@dataclass
class ExtractedText:
    file_id: str
    name: str = ""
    mime: str = ""
    modified_time: str = ""
    text: str = ""
    char_count: int = 0
    truncated: bool = False
    extracted_at: str = ""
    injection_hits: list[str] = field(default_factory=list)

    @property
    def suspicious(self) -> bool:
        return bool(self.injection_hits)

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)

    @classmethod
    def from_dict(cls, data: dict[str, Any]) -> "ExtractedText":
        return cls(**{k: v for k, v in data.items() if k in cls.__dataclass_fields__})


# ---------------------------------------------------------------------------
# Cache
# ---------------------------------------------------------------------------

class TextCache:
    """Cache text đã trích, khoá theo (file_id, modified_time)."""

    def __init__(self, root: str | Path = "state/cache") -> None:
        self.root = Path(root)

    @staticmethod
    def key(file_id: str, modified_time: str) -> str:
        digest = hashlib.sha256(f"{file_id}|{modified_time}".encode()).hexdigest()
        return digest[:32]

    def _path(self, file_id: str, modified_time: str) -> Path:
        return self.root / f"{self.key(file_id, modified_time)}.json"

    def get(self, file_id: str, modified_time: str) -> ExtractedText | None:
        path = self._path(file_id, modified_time)
        if not path.exists():
            return None
        try:
            return ExtractedText.from_dict(json.loads(path.read_text(encoding="utf-8")))
        except (json.JSONDecodeError, TypeError, ValueError):
            # Cache hỏng thì bỏ qua và trích lại — không bao giờ để cache làm sập luồng
            return None

    def put(self, extracted: ExtractedText) -> None:
        self.root.mkdir(parents=True, exist_ok=True)
        path = self._path(extracted.file_id, extracted.modified_time)
        path.write_text(json.dumps(extracted.to_dict(), ensure_ascii=False),
                        encoding="utf-8")


# ---------------------------------------------------------------------------
# Các bộ trích xuất theo định dạng
# ---------------------------------------------------------------------------

def _extract_plain(data: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "cp1258", "latin-1"):
        try:
            return data.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    return data.decode("utf-8", errors="replace")


def _extract_docx(data: bytes) -> str:
    try:
        import docx  # noqa: PLC0415
    except ImportError as exc:  # pragma: no cover - phụ thuộc tuỳ chọn
        raise ExtractionError("Thiếu python-docx — chạy: pip install python-docx") from exc

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_xlsx(data: bytes) -> str:
    try:
        import openpyxl  # noqa: PLC0415
    except ImportError as exc:  # pragma: no cover - phụ thuộc tuỳ chọn
        raise ExtractionError("Thiếu openpyxl — chạy: pip install openpyxl") from exc

    workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    parts: list[str] = []
    for sheet in workbook.worksheets:
        parts.append(f"### {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None and str(c).strip()]
            if cells:
                parts.append(" | ".join(cells))
    workbook.close()
    return "\n".join(parts)


def _extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader  # noqa: PLC0415
    except ImportError as exc:  # pragma: no cover - phụ thuộc tuỳ chọn
        raise ExtractionError("Thiếu pypdf — chạy: pip install pypdf") from exc

    reader = PdfReader(io.BytesIO(data))
    pages = [(page.extract_text() or "").strip() for page in reader.pages]
    text = "\n\n".join(p for p in pages if p)
    if not text.strip():
        raise ExtractionError(
            "PDF không có lớp text — nhiều khả năng là bản scan, cần OCR."
        )
    return text


_EXTRACTORS = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": _extract_docx,
    "application/msword": _extract_docx,
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": _extract_xlsx,
    "application/vnd.ms-excel": _extract_xlsx,
    "application/pdf": _extract_pdf,
    "text/plain": _extract_plain,
    "text/markdown": _extract_plain,
    "text/csv": _extract_plain,
}


def extract_bytes(data: bytes, mime: str) -> str:
    """Trích text thô theo định dạng. Không cache, không dò injection."""
    extractor = _EXTRACTORS.get(mime)
    if extractor is None:
        raise ExtractionError(f"Chưa hỗ trợ định dạng {mime!r}.")
    return extractor(data)


# ---------------------------------------------------------------------------
# Tải + trích + cache
# ---------------------------------------------------------------------------

class Downloader(Protocol):
    """Chỉ có phép tải về — không có phương thức nào ghi hay xoá."""

    def download(self, file_id: str, mime: str) -> bytes:
        ...


def extract_file(
    file_id: str,
    mime: str,
    modified_time: str,
    downloader: Downloader,
    cache: TextCache | None = None,
    name: str = "",
    max_chars: int = MAX_CHARS,
) -> ExtractedText:
    """Lấy text của một tài liệu, ưu tiên cache."""
    if cache is not None:
        cached = cache.get(file_id, modified_time)
        if cached is not None:
            return cached

    effective_mime = GOOGLE_EXPORT_MIMES.get(mime, mime)
    data = downloader.download(file_id, effective_mime)
    text = extract_bytes(data, effective_mime).strip()

    truncated = len(text) > max_chars
    if truncated:
        text = text[:max_chars]

    hits: Sequence[InjectionHit] = scan_for_injection(text)
    extracted = ExtractedText(
        file_id=file_id,
        name=name,
        mime=mime,
        modified_time=modified_time,
        text=text,
        char_count=len(text),
        truncated=truncated,
        extracted_at=datetime.now(timezone.utc).isoformat(timespec="seconds"),
        injection_hits=[str(h) for h in hits],
    )

    if cache is not None:
        cache.put(extracted)
    return extracted


# ---------------------------------------------------------------------------
# Chọn trích đoạn liên quan
# ---------------------------------------------------------------------------

def _normalize(text: str) -> str:
    decomposed = unicodedata.normalize("NFD", text.lower())
    stripped = "".join(c for c in decomposed if unicodedata.category(c) != "Mn")
    return stripped.replace("đ", "d")


def find_excerpts(text: str, keywords: Sequence[str], window: int = 400,
                  max_excerpts: int = 5) -> list[str]:
    """Cắt ra vài đoạn quanh từ khoá thay vì đưa cả tài liệu vào prompt.

    Một bảng giá 20 trang mà câu hỏi chỉ liên quan tới ba dòng thì đưa cả file
    vào vừa tốn tiền vừa làm loãng tín hiệu. Không tìm thấy từ khoá nào thì trả
    về phần đầu tài liệu — trung thực hơn là trả về rỗng.
    """
    if not text:
        return []

    haystack = _normalize(text)
    spans: list[tuple[int, int]] = []
    for keyword in keywords:
        needle = _normalize(keyword).strip()
        if not needle:
            continue
        for match in re.finditer(re.escape(needle), haystack):
            start = max(0, match.start() - window // 2)
            end = min(len(text), match.end() + window // 2)
            spans.append((start, end))
            if len(spans) >= max_excerpts * 4:
                break

    if not spans:
        return [text[: window * 2].strip()]

    # Gộp các cửa sổ chồng nhau để không lặp lại cùng một đoạn
    spans.sort()
    merged: list[list[int]] = [list(spans[0])]
    for start, end in spans[1:]:
        if start <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], end)
        else:
            merged.append([start, end])

    return [text[start:end].strip() for start, end in merged[:max_excerpts]]


# ---------------------------------------------------------------------------
# Bộ tải Drive thật
# ---------------------------------------------------------------------------

class GoogleDriveDownloader:
    """Tải file bằng service account CHỈ ĐỌC.

    Import SDK nằm trong hàm khởi tạo để phần logic ở trên chạy và test được
    mà không cần cài SDK Google.
    """

    SCOPES = ["https://www.googleapis.com/auth/drive.readonly"]

    def __init__(self, credentials_path: str) -> None:
        from google.oauth2 import service_account  # noqa: PLC0415
        from googleapiclient.discovery import build  # noqa: PLC0415

        creds = service_account.Credentials.from_service_account_file(
            credentials_path, scopes=self.SCOPES
        )
        self._service = build("drive", "v3", credentials=creds, cache_discovery=False)

    def download(self, file_id: str, mime: str) -> bytes:
        from googleapiclient.http import MediaIoBaseDownload  # noqa: PLC0415

        files = self._service.files()
        # File Google Docs gốc phải export, file thường thì tải trực tiếp
        if mime in GOOGLE_EXPORT_MIMES.values() and self._is_native(file_id):
            request = files.export_media(fileId=file_id, mimeType=mime)
        else:
            request = files.get_media(fileId=file_id, supportsAllDrives=True)

        buffer = io.BytesIO()
        downloader = MediaIoBaseDownload(buffer, request)
        done = False
        while not done:
            _, done = downloader.next_chunk()
        return buffer.getvalue()

    def _is_native(self, file_id: str) -> bool:
        meta = self._service.files().get(
            fileId=file_id, fields="mimeType", supportsAllDrives=True
        ).execute()
        return str(meta.get("mimeType", "")).startswith("application/vnd.google-apps.")
